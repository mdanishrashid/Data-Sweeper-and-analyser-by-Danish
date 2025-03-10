import streamlit as st
import pandas as pd
import os
from io import BytesIO
from docx import Document
from fpdf import FPDF
from pdf2docx import Converter

def convert_csv_to_excel(csv_file):
    st.info("📂 Converting CSV to Excel...")
    df = pd.read_csv(csv_file)
    output = BytesIO()
    df.to_excel(output, index=False, engine='openpyxl')
    output.seek(0)
    return output

def convert_excel_to_csv(excel_file):
    st.info("📂 Converting Excel to CSV...")
    df = pd.read_excel(excel_file, engine='openpyxl')
    output = BytesIO()
    df.to_csv(output, index=False)
    output.seek(0)
    return output

def convert_word_to_pdf(word_file):
    st.info("📄 Converting Word to PDF...")
    doc = Document(word_file)
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    
    for para in doc.paragraphs:
        pdf.cell(200, 10, txt=para.text, ln=True)
    
    output = BytesIO()
    pdf.output(output)
    output.seek(0)
    return output

def convert_pdf_to_word(pdf_file):
    st.info("📄 Converting PDF to Word...")
    output = BytesIO()
    with open("temp.pdf", "wb") as temp_pdf:
        temp_pdf.write(pdf_file.getvalue())
    
    docx_filename = "converted.docx"
    cv = Converter("temp.pdf")
    cv.convert(docx_filename, start=0, end=None)
    cv.close()
    
    with open(docx_filename, "rb") as docx_file:
        output.write(docx_file.read())
    os.remove("temp.pdf")
    os.remove(docx_filename)
    output.seek(0)
    return output

def clean_data(df):
    st.info("🧹 Cleaning Data: Removing duplicates & missing values...")
    df.dropna(inplace=True)
    df.drop_duplicates(inplace=True)
    return df

def analyze_data(df):
    st.success("📊 Data Analysis Completed!")
    st.write("### 🔍 Data Summary")
    st.write(df.describe())
    st.write("### ❓ Missing Values")
    st.write(df.isnull().sum())

def main():
    st.title("🧹 Data Sweeper & Analyzer 📊")
    st.sidebar.title("🔍 Navigation")
    option = st.sidebar.radio("📌 Choose an option", ["CSV ↔ Excel", "Word ↔ PDF", "Data Cleaning & Analysis"])
    
    if option == "CSV ↔ Excel":
        file = st.file_uploader("📂 Upload CSV/Excel File", type=["csv", "xlsx"])
        if file:
            file_type = file.name.split(".")[-1]
            if file_type == "csv":
                output = convert_csv_to_excel(file)
                st.download_button("⬇️ Download Excel", output, file_name="converted.xlsx")
            elif file_type == "xlsx":
                output = convert_excel_to_csv(file)
                st.download_button("⬇️ Download CSV", output, file_name="converted.csv")
    
    elif option == "Word ↔ PDF":
        file = st.file_uploader("📄 Upload Word/PDF File", type=["docx", "pdf"])
        if file:
            file_type = file.name.split(".")[-1]
            if file_type == "docx":
                output = convert_word_to_pdf(file)
                st.download_button("⬇️ Download PDF", output, file_name="converted.pdf")
            elif file_type == "pdf":
                output = convert_pdf_to_word(file)
                st.download_button("⬇️ Download Word", output, file_name="converted.docx")
    
    elif option == "Data Cleaning & Analysis":
        file = st.file_uploader("📊 Upload CSV File", type=["csv"])
        if file:
            df = pd.read_csv(file)
            df = clean_data(df)
            analyze_data(df)
            output = BytesIO()
            df.to_csv(output, index=False)
            output.seek(0)
            st.download_button("⬇️ Download Cleaned Data", output, file_name="cleaned_data.csv")
    
if __name__ == "__main__":
    main()
